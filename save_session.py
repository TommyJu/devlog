from pathlib import Path
from datetime import timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font



DEVLOG_PATH = Path("devlog.xlsx")

SESSION_HEADERS = [
    "Start Date",
    "End Date",
    "Time Elapsed",
    "Session Notes"
]

def initialize_sessions(worksheet):
    """Initialize the Sessions worksheet."""

    worksheet.append(SESSION_HEADERS)

    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 30
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 200

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    worksheet.freeze_panes = "A2"


def initialize_stats(worksheet):
    """Initialize the Stats worksheet."""

    worksheet["A1"] = "Dev Log Statistics"
    worksheet["A1"].font = Font(bold=True, size=16)

    worksheet["A3"] = "Period"
    worksheet["B3"] = "Total Time"

    for cell in worksheet[3]:
        cell.font = Font(bold=True)

    worksheet["A4"] = "Today"
    worksheet["A5"] = "This Week"
    worksheet["A6"] = "This Month"
    worksheet["A7"] = "All Time"

    worksheet["B4"] = (
        '=SUMIFS(Sessions!C:C,'
        'Sessions!A:A,">="&TODAY(),'
        'Sessions!A:A,"<"&TODAY()+1)'
    )

    worksheet["B5"] = (
        '=SUMIFS(Sessions!C:C,'
        'Sessions!A:A,">="&TODAY()-WEEKDAY(TODAY(),2)+1,'
        'Sessions!A:A,"<"&TODAY()-WEEKDAY(TODAY(),2)+8)'
    )

    worksheet["B6"] = (
        '=SUMIFS(Sessions!C:C,'
        'Sessions!A:A,">="&EOMONTH(TODAY(),-1)+1,'
        'Sessions!A:A,"<"&EOMONTH(TODAY(),0)+1)'
    )

    worksheet["B7"] = "=SUM(Sessions!C:C)"

    for row in range(4, 8):
        worksheet[f"B{row}"].number_format = "[h]:mm:ss"

    worksheet.column_dimensions["A"].width = 20
    worksheet.column_dimensions["B"].width = 20


def initialize_workbook():
    """Load or create the dev log workbook and its worksheets."""

    if DEVLOG_PATH.exists():
        workbook = load_workbook(DEVLOG_PATH)
    else:
        workbook = Workbook()

    # Get or create Sessions worksheet.
    if "Sessions" in workbook.sheetnames:
        sessions = workbook["Sessions"]
    else:
        sessions = workbook.create_sheet("Sessions")

        # Remove the default blank worksheet if appropriate.
        if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) > 1:
            del workbook["Sheet"]

        initialize_sessions(sessions)

    # Get or create Stats worksheet.
    if "Stats" in workbook.sheetnames:
        stats = workbook["Stats"]
    else:
        stats = workbook.create_sheet("Stats")
        initialize_stats(stats)

    return workbook, sessions, stats


def save_session(app_state):
    """Append a completed session to the dev log."""

    timer = app_state.timer

    workbook, sessions, _ = initialize_workbook()

    sessions.append([
        timer.start_datetime,
        timer.end_datetime,
        timedelta(seconds=timer.elapsed_seconds),
        app_state.session_notes
    ])

    # Format the new row.
    row = sessions.max_row

    sessions.cell(row=row, column=1).number_format = "yyyy-mm-dd hh:mm AM/PM"
    sessions.cell(row=row, column=2).number_format = "yyyy-mm-dd hh:mm AM/PM"
    sessions.cell(row=row, column=3).number_format = "[h]:mm:ss"

    workbook.save(DEVLOG_PATH)